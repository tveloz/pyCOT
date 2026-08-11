"""
╔══════════════════════════════════════════════════════════════════════════════╗
║  build_catalogue.py — Scan all reaction networks and write a catalogue     ║
╚══════════════════════════════════════════════════════════════════════════════╝

WHAT IT DOES
------------
Walks the entire data/biomodels/ directory tree, loads every .txt reaction
network using pyCOT's read_txt parser, counts species and reactions, and
writes a summary catalogue file.

WHAT IT OUTPUTS
---------------
  • network_catalogue.csv   — always written (one row per network, sorted by
                               reaction count ascending).
  • network_catalogue.xlsx  — written additionally when openpyxl is installed
                               and CSV_ONLY = False. Includes:
                               - Colour-coded reaction count column
                                 (green ≤100 | yellow ≤500 | orange ≤2000 | red >2000)
                               - Frozen header row and auto-filter.

  Both files are written to the COT_fundamental_Generators/ project folder.
  The CSV can be passed to compare_oracles.py and run_network.py via the
  CATALOGUE_CSV setting to speed up network discovery.

HOW TO RUN
----------
  1. (Optional) edit the CONFIGURATION section below.
  2. Press ▶ (play) in VS Code, or run:
       python projects/COT_fundamental_Generators/scripts/build_catalogue.py
"""

# ── Path setup (do not edit) ──────────────────────────────────────────────────
from __future__ import annotations
import csv, os, sys, time
_here = os.path.dirname(os.path.abspath(__file__))
_proj = os.path.normpath(os.path.join(_here, ".."))
_repo = os.path.normpath(os.path.join(_here, "..", "..", ".."))
for _p in (_proj, os.path.join(_repo, "src")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

# ── Imports (do not edit) ─────────────────────────────────────────────────────
from pyCOT.io.functions import read_txt

# ╔══════════════════════════════════════════════════════════════════════════════╗
# ║  CONFIGURATION — edit this block, then press Play                          ║
# ╚══════════════════════════════════════════════════════════════════════════════╝

# ── Output format ─────────────────────────────────────────────────────────────
# False → write both CSV and Excel (.xlsx) if openpyxl is installed
# True  → write only CSV (skip Excel even if openpyxl is available)
CSV_ONLY = False

# ╔══════════════════════════════════════════════════════════════════════════════╗
# ║  Script body — no need to edit below this line                             ║
# ╚══════════════════════════════════════════════════════════════════════════════╝

_DATA_ROOT = os.path.join(_repo, "data", "biomodels")
_OUT_DIR   = _proj   # projects/COT_fundamental_Generators/

# ── Discover all .txt files ───────────────────────────────────────────────────
def _discover() -> list[tuple[str, str, str]]:
    """
    Return one (short_name, collection_folder, abs_path) per unique network name.
    When the same file exists in multiple directories, keep the shortest path
    (typically the most canonical flat copy).
    """
    # name -> (folder, path, path_len)
    seen: dict[str, tuple[str, str, int]] = {}
    for root, _dirs, files in os.walk(_DATA_ROOT):
        for fname in sorted(files):
            if not fname.endswith(".txt"):
                continue
            full   = os.path.join(root, fname)
            name   = os.path.splitext(fname)[0]
            if name.startswith("bigg_"):
                name = name[5:]
            folder = os.path.basename(root)
            path_len = len(full)
            if name not in seen or path_len < seen[name][2]:
                seen[name] = (folder, full, path_len)
    return [(name, folder, path) for name, (folder, path, _) in sorted(seen.items())]

# ── Count species and reactions via pyCOT ────────────────────────────────────
def _count(path: str) -> tuple[int, int]:
    try:
        rn    = read_txt(path)
        n_rxn = len(rn.reactions())   # reactions() is a method, not an attribute
        n_sp  = len(rn.species())     # species()  is a method, not an attribute
        return n_sp, n_rxn
    except Exception as e:
        return -1, -1

# ── Main scan ─────────────────────────────────────────────────────────────────
files = _discover()
print(f"Found {len(files)} network files. Scanning...\n")

rows: list[dict] = []
t_start = time.perf_counter()

for i, (name, folder, path) in enumerate(files):
    n_sp, n_rxn = _count(path)
    rel = os.path.relpath(path, _repo)
    rows.append({"name": name, "collection": folder,
                 "species": n_sp, "reactions": n_rxn, "rel_path": rel})
    if (i + 1) % 20 == 0 or i == len(files) - 1:
        elapsed = time.perf_counter() - t_start
        print(f"  {i+1}/{len(files)}  ({elapsed:.1f}s)")

rows.sort(key=lambda r: (r["reactions"] < 0, r["reactions"]))

# ── Write CSV ─────────────────────────────────────────────────────────────────
csv_path = os.path.join(_OUT_DIR, "network_catalogue.csv")
with open(csv_path, "w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(
        f, fieldnames=["name", "collection", "species", "reactions", "rel_path"]
    )
    writer.writeheader()
    writer.writerows(rows)
print(f"\nCSV  : {csv_path}")

# ── Write Excel ───────────────────────────────────────────────────────────────
xlsx_path = os.path.join(_OUT_DIR, "network_catalogue.xlsx")

if not CSV_ONLY:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils  import get_column_letter   # noqa: F401

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Networks"

        hdr_font = Font(bold=True, color="FFFFFF")
        hdr_fill = PatternFill(fill_type="solid", fgColor="1F4E79")
        thin     = Side(style="thin", color="AAAAAA")
        thin_bdr = Border(left=thin, right=thin, top=thin, bottom=thin)

        headers = ["Name", "Collection", "Species", "Reactions", "Relative path"]
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font      = hdr_font
            cell.fill      = hdr_fill
            cell.border    = thin_bdr
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 20

        for row_i, r in enumerate(rows, start=2):
            rxn = r["reactions"]
            if   rxn < 0:    size_color = "F4CCCC"   # unknown → red
            elif rxn <= 100: size_color = "D9EAD3"   # small   → green
            elif rxn <= 500: size_color = "FFF3CC"   # medium  → yellow
            elif rxn <= 2000:size_color = "FCE5CD"   # large   → orange
            else:            size_color = "F4CCCC"   # very large → red

            vals = [r["name"], r["collection"], r["species"], r["reactions"], r["rel_path"]]
            for col, v in enumerate(vals, start=1):
                cell = ws.cell(row=row_i, column=col, value=v)
                cell.border = thin_bdr
                if col == 4:
                    cell.fill = PatternFill(fill_type="solid", fgColor=size_color)
                elif row_i % 2 == 0:
                    cell.fill = PatternFill(fill_type="solid", fgColor="F0F4FA")

        ws.column_dimensions["A"].width = 36
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions["E"].width = 60
        ws.freeze_panes       = "A2"
        ws.auto_filter.ref    = f"A1:E{len(rows)+1}"

        wb.save(xlsx_path)
        print(f"Excel: {xlsx_path}")

    except ImportError:
        print("openpyxl not installed — Excel skipped.  Install with: pip install openpyxl")

print(f"\nDone. {len(rows)} networks catalogued in {time.perf_counter()-t_start:.1f}s.")
